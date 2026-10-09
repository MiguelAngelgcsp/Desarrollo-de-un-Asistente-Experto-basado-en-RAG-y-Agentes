"""
Paso 1 — Ingesta.
Extrae el texto de PDF, DOCX, TXT y HTML y lo convierte en objetos Document de
LangChain con metadatos para poder citar la fuente:
    source  -> nombre del archivo
    pagina  -> número de página (1-based; solo PDF)
    tipo    -> extensión del archivo
"""

import re
from pathlib import Path

from langchain_core.documents import Document

from src.config import DOCS_DIR, EXTENSIONES_SOPORTADAS, PATRONES_RUIDO

_RUIDO = [re.compile(p, re.MULTILINE) for p in PATRONES_RUIDO]

def limpiar_texto(texto: str) -> str:
    """Quita encabezados/pies repetidos y normaliza espacios, conservando los párrafos."""
    for patron in _RUIDO:
        texto = patron.sub("", texto)
    texto = texto.replace("\u00a0", " ")
    texto = re.sub(r"[ \t]+", " ", texto)
    texto = re.sub(r" *\n *", "\n", texto)
    texto = re.sub(r"\n{3,}", "\n\n", texto)
    return texto.strip()


def _cargar_pdf(ruta: Path):
    try:
        import pymupdf  # mejor extracción que pypdf (espacios y orden de lectura)

        with pymupdf.open(ruta) as pdf:
            textos = [p.get_text() for p in pdf]
    except ImportError:
        from pypdf import PdfReader

        textos = [(p.extract_text() or "") for p in PdfReader(str(ruta)).pages]
    return [(i + 1, t) for i, t in enumerate(textos)]


def _cargar_docx(ruta: Path):
    from docx import Document as DocxDocument

    doc = DocxDocument(str(ruta))
    partes = [p.text for p in doc.paragraphs if p.text.strip()]
    for tabla in doc.tables:
        for fila in tabla.rows:
            partes.append(" | ".join(c.text.strip() for c in fila.cells))
    return [(None, "\n\n".join(partes))]


def _cargar_txt(ruta: Path):
    return [(None, ruta.read_text(encoding="utf-8", errors="ignore"))]


def _cargar_html(ruta: Path):
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(ruta.read_text(encoding="utf-8", errors="ignore"), "html.parser")
    for etiqueta in soup(["script", "style", "nav", "footer", "header"]):
        etiqueta.decompose()
    return [(None, soup.get_text("\n"))]


_CARGADORES = {
    ".pdf": _cargar_pdf,
    ".docx": _cargar_docx,
    ".txt": _cargar_txt,
    ".html": _cargar_html,
    ".htm": _cargar_html,
}


def load_documents(docs_dir: Path = DOCS_DIR):
    if not docs_dir.exists():
        raise FileNotFoundError(f"No existe la carpeta del corpus: '{docs_dir}'")

    archivos = sorted(
        f for f in docs_dir.iterdir() if f.is_file() and f.suffix.lower() in EXTENSIONES_SOPORTADAS
    )
    if not archivos:
        raise FileNotFoundError(
            f"No hay archivos {sorted(EXTENSIONES_SOPORTADAS)} en '{docs_dir}'."
        )

    print(f"Corpus en '{docs_dir}' ({len(archivos)} archivos):")
    documentos = []
    for archivo in archivos:
        ext = archivo.suffix.lower()
        unidades = _CARGADORES[ext](archivo)
        cargadas = 0
        for pagina, texto in unidades:
            texto = limpiar_texto(texto)
            if not texto:
                continue
            metadata = {"source": archivo.name, "tipo": ext.lstrip(".")}
            if pagina is not None:
                metadata["pagina"] = pagina
            documentos.append(Document(page_content=texto, metadata=metadata))
            cargadas += 1
        print(f"  - {archivo.name}: {cargadas} unidades de texto ({ext})")

    print(f"Total de unidades cargadas: {len(documentos)}")
    return documentos
